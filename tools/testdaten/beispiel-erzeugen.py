#!/usr/bin/env python3
"""
Erzeugt tools/testdaten/beispiel.pptx fuer den Browsertest des PowerPoint-Imports.

Die Datei ist wie eine echte PowerPoint gepackt (Deflate) und nutzt die Vorlage von
python-pptx (4:3) mit Platzhaltern, Design-Farben und Listenstilen aus Master und
Layout. Dazu kommen Teile, die python-pptx nicht direkt schreibt (Animation,
Uebergang in mc:AlternateContent), als XML.

Aufruf: python3 tools/testdaten/beispiel-erzeugen.py   (braucht: pip install python-pptx)
"""
import io
import os
import struct
import zlib

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

ORDNER = os.path.dirname(os.path.abspath(__file__))
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def png(breite, hoehe, farbe):
    """Einfarbiges PNG ohne Pillow"""
    zeile = b"\x00" + bytes(farbe) * breite
    roh = zlib.compress(zeile * hoehe)

    def block(art, daten):
        return struct.pack(">I", len(daten)) + art + daten + struct.pack(">I", zlib.crc32(art + daten))

    return b"\x89PNG\r\n\x1a\n" + block(b"IHDR", struct.pack(">IIBBBBB", breite, hoehe, 8, 2, 0, 0, 0)) + block(b"IDAT", roh) + block(b"IEND", b"")


def xml(text):
    return etree.fromstring(text)


prs = Presentation()

# 1 Titelfolie: dunkler Hintergrund, Notizen, Uebergang (modern mit Ersatz)
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "Der Wasserkreislauf"
s1.placeholders[1].text = "Klasse 5 · Geografie"
s1.background.fill.solid()
s1.background.fill.fore_color.rgb = RGBColor(0x21, 0x29, 0x5C)
for p in s1.shapes.title.text_frame.paragraphs:
    for r in p.runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
s1.notes_slide.notes_text_frame.text = "Begrüssung und Vorwissen sammeln."
s1._element.append(
    xml(
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
        '<mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="1500"><p14:vortex dir="r"/></p:transition></mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
    )
)

# 2 Titel und Inhalt: Aufzaehlung mit zweiter Stufe, Stern mit Animation, Bild
s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = "Stationen"
koerper = s2.placeholders[1].text_frame
koerper.text = "Verdunstung"
for text, stufe in [("Kondensation", 0), ("Wolken entstehen", 1), ("Niederschlag", 0)]:
    p = koerper.add_paragraph()
    p.text = text
    p.level = stufe
stern = s2.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(6.5), Inches(4.5), Inches(2.5), Inches(2.2))
stern.fill.solid()
stern.fill.fore_color.rgb = RGBColor(0xE4, 0x57, 0x2E)
stern.text_frame.text = "Merke!"
s2.shapes.add_picture(io.BytesIO(png(40, 30, (0x06, 0x5A, 0x82))), Inches(6.5), Inches(1.8), Inches(2.4), Inches(1.8))
s2._element.append(
    xml(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="5" presetID="2" presetClass="entr" presetSubtype="4" fill="hold" nodeType="clickEffect">'
        '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        '<p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
        '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
        '<p:to><p:strVal val="visible"/></p:to></p:set>'
        "</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
        "</p:childTnLst></p:cTn><p:prevCondLst/><p:nextCondLst/></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>" % stern.shape_id
    )
)

# 3 Nur Titel: Tabelle, Linie, Gruppe
s3 = prs.slides.add_slide(prs.slide_layouts[5])
s3.shapes.title.text = "Messwerte"
tabelle = s3.shapes.add_table(3, 3, Inches(0.5), Inches(1.8), Inches(6), Inches(1.5)).table
for z, zeile in enumerate([["Ort", "Regen", "Sonne"], ["Bern", "1000 mm", "1700 h"], ["Lugano", "1600 mm", "2100 h"]]):
    for sp, wert in enumerate(zeile):
        tabelle.cell(z, sp).text = wert
linie = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(4), Inches(9.5), Inches(4))
linie.line.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
linie.line.width = Pt(4)
gruppe = s3.shapes.add_group_shape()
a = gruppe.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(4.5), Inches(2), Inches(2))
a.text_frame.text = "A"
b = gruppe.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(4.5), Inches(2), Inches(2))
b.text_frame.text = "B"

# 4 Leer, ausgeblendet: Textfeld mit eigener Schrift, Diagramm, Video
s4 = prs.slides.add_slide(prs.slide_layouts[6])
s4._element.set("show", "0")
feld = s4.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.2))
lauf = feld.text_frame.paragraphs[0].add_run()
lauf.text = "Zusatz für Schnelle"
lauf.font.name = "Georgia"
lauf.font.size = Pt(40)
lauf.font.bold = True
lauf.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)
daten = CategoryChartData()
daten.categories = ["Jan", "Feb"]
daten.add_series("Regen", (60, 55))
s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2), Inches(4), Inches(3), daten)
s4.shapes.add_movie(io.BytesIO(b"\x00\x00\x00\x18ftypmp42"), Inches(5), Inches(2), Inches(4), Inches(3), mime_type="video/mp4", poster_frame_image=io.BytesIO(png(16, 9, (0, 0, 0))))

prs.core_properties.title = "Wasserkreislauf"
prs.save(os.path.join(ORDNER, "beispiel.pptx"))
print("geschrieben:", os.path.join(ORDNER, "beispiel.pptx"))
